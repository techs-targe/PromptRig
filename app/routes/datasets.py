"""Dataset management API endpoints.

Based on specification in docs/req.txt section 4.6 (データセット)
Phase 2 implementation.
"""

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session
from sqlalchemy import text
from typing import List, Dict, Any, Optional
from pydantic import BaseModel
from datetime import datetime
import os
import tempfile
import csv
import io
import json

from backend.database import get_db, Dataset, Job, JobItem, Project, ProjectDataset
from backend.dataset import DatasetImporter

router = APIRouter()


class DatasetResponse(BaseModel):
    """Dataset response model."""
    id: int
    project_id: int  # Original owner (backward compatibility)
    name: str
    source_file_name: str
    sqlite_table_name: str
    created_at: str
    row_count: int = 0
    project_ids: List[int] = []  # All associated project IDs (many-to-many)


class DatasetPreviewResponse(BaseModel):
    """Dataset preview response model."""
    dataset_id: int
    name: str
    columns: List[str]
    rows: List[Dict[str, Any]]
    total_count: int


@router.get("/api/datasets", response_model=List[DatasetResponse])
def list_datasets(project_id: int = None, include_associated: bool = True, db: Session = Depends(get_db)):
    """List datasets, optionally filtered by project.

    Specification: docs/req.txt section 4.6.1
    Phase 2

    Args:
        project_id: Filter by project ID (optional)
        include_associated: If True, also include datasets associated via ProjectDataset (not just owned)
    """
    if project_id:
        if include_associated:
            # Get datasets owned by project OR associated via ProjectDataset
            owned_datasets = db.query(Dataset).filter(Dataset.project_id == project_id).all()
            associated_dataset_ids = db.query(ProjectDataset.dataset_id).filter(
                ProjectDataset.project_id == project_id
            ).all()
            associated_ids = [a[0] for a in associated_dataset_ids]

            # Combine unique datasets
            all_dataset_ids = set([d.id for d in owned_datasets] + associated_ids)
            datasets = db.query(Dataset).filter(Dataset.id.in_(all_dataset_ids)).order_by(Dataset.created_at.desc()).all()
        else:
            datasets = db.query(Dataset).filter(Dataset.project_id == project_id).order_by(Dataset.created_at.desc()).all()
    else:
        datasets = db.query(Dataset).order_by(Dataset.created_at.desc()).all()

    result = []
    for dataset in datasets:
        # Get row count from importer
        importer = DatasetImporter(db)
        row_count = importer._get_row_count(dataset.sqlite_table_name)

        # Get all associated project IDs
        associations = db.query(ProjectDataset.project_id).filter(
            ProjectDataset.dataset_id == dataset.id
        ).all()
        associated_project_ids = [a[0] for a in associations]

        # Include original owner in the list
        all_project_ids = list(set([dataset.project_id] + associated_project_ids))

        result.append(DatasetResponse(
            id=dataset.id,
            project_id=dataset.project_id,
            name=dataset.name,
            source_file_name=dataset.source_file_name,
            sqlite_table_name=dataset.sqlite_table_name,
            created_at=dataset.created_at,
            row_count=row_count,
            project_ids=all_project_ids
        ))

    return result


# ========== Hugging Face Dataset Endpoints ==========
# NOTE: These must be defined before /api/datasets/{dataset_id} routes
# to avoid path matching conflicts


def _check_huggingface_enabled():
    """Check if HuggingFace import feature is enabled.

    Raises HTTPException 403 if the feature is disabled.
    """
    env_value = os.getenv("HUGGINGFACE_IMPORT_ENABLED", "false").lower()
    if env_value not in ("true", "1", "yes", "on"):
        raise HTTPException(
            status_code=403,
            detail="Hugging Face import feature is disabled. Set HUGGINGFACE_IMPORT_ENABLED=true in .env to enable."
        )


class HuggingFaceSearchResult(BaseModel):
    """Single search result for Hugging Face dataset."""
    id: str  # Full dataset ID (e.g., "username/dataset-name")
    name: str  # Short name
    author: str
    description: str
    downloads: int
    likes: int
    tags: List[str]
    is_gated: bool
    last_modified: Optional[str] = None


class HuggingFaceSearchResponse(BaseModel):
    """Response model for Hugging Face dataset search."""
    query: str
    count: int
    results: List[HuggingFaceSearchResult]


class HuggingFaceDatasetInfoResponse(BaseModel):
    """Response model for Hugging Face dataset info."""
    name: str
    description: str
    splits: List[str]
    features: Dict[str, str]
    size_info: Dict[str, Dict[str, Any]]
    is_gated: bool
    requires_auth: bool
    warning: Optional[str] = None


class HuggingFacePreviewResponse(BaseModel):
    """Response model for Hugging Face dataset preview."""
    name: str
    split: str
    columns: List[str]
    rows: List[Dict[str, Any]]
    total_count: int


class HuggingFaceImportRequest(BaseModel):
    """Request model for Hugging Face dataset import."""
    project_id: int
    dataset_name: str
    split: str
    display_name: str
    row_limit: Optional[int] = None
    columns: Optional[List[str]] = None
    add_row_id: bool = False


@router.get("/api/datasets/huggingface/search", response_model=HuggingFaceSearchResponse)
def search_huggingface_datasets(
    query: str,
    limit: int = 20,
    db: Session = Depends(get_db)
):
    """Search for datasets on Hugging Face Hub by keyword.

    Args:
        query: Search keyword (e.g., "question answering", "sentiment analysis")
        limit: Maximum number of results (default 20, max 100)

    Returns:
        List of matching datasets with metadata
    """
    _check_huggingface_enabled()

    from backend.dataset import HuggingFaceImporter

    if not query or not query.strip():
        raise HTTPException(status_code=400, detail="Search query is required")

    try:
        importer = HuggingFaceImporter(db)
        results = importer.search_datasets(query.strip(), limit=limit)

        return HuggingFaceSearchResponse(
            query=query.strip(),
            count=len(results),
            results=[
                HuggingFaceSearchResult(
                    id=r["id"],
                    name=r["name"],
                    author=r["author"],
                    description=r["description"][:200] if r["description"] else "",
                    downloads=r["downloads"],
                    likes=r["likes"],
                    tags=r["tags"][:5] if r["tags"] else [],  # Limit tags
                    is_gated=r["is_gated"],
                    last_modified=r["last_modified"]
                )
                for r in results
            ]
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to search datasets: {str(e)}")


@router.get("/api/datasets/huggingface/info", response_model=HuggingFaceDatasetInfoResponse)
def get_huggingface_dataset_info(
    name: str,
    db: Session = Depends(get_db)
):
    """Get information about a Hugging Face dataset.

    Args:
        name: Dataset name (e.g., "squad", "imdb", "username/dataset")

    Returns:
        Dataset metadata including splits, features, and size info
    """
    _check_huggingface_enabled()

    from backend.dataset import HuggingFaceImporter

    try:
        importer = HuggingFaceImporter(db)
        info = importer.get_dataset_info(name)

        return HuggingFaceDatasetInfoResponse(
            name=info.name,
            description=info.description,
            splits=info.splits,
            features=info.features,
            size_info=info.size_info,
            is_gated=info.is_gated,
            requires_auth=info.requires_auth,
            warning=info.warning
        )
    except ValueError as e:
        if "not found" in str(e).lower():
            raise HTTPException(status_code=404, detail=str(e))
        elif "authentication" in str(e).lower() or "token" in str(e).lower():
            raise HTTPException(status_code=401, detail=str(e))
        else:
            raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to get dataset info: {str(e)}")


@router.get("/api/datasets/huggingface/preview", response_model=HuggingFacePreviewResponse)
def preview_huggingface_dataset(
    name: str,
    split: str,
    limit: int = 10,
    db: Session = Depends(get_db)
):
    """Get a preview of a Hugging Face dataset.

    Args:
        name: Dataset name
        split: Split to preview (e.g., "train", "validation")
        limit: Maximum number of rows to preview (default: 10)

    Returns:
        Preview with columns, rows, and total count
    """
    _check_huggingface_enabled()

    from backend.dataset import HuggingFaceImporter

    if limit > 100:
        limit = 100  # Cap preview limit

    try:
        importer = HuggingFaceImporter(db)
        preview = importer.get_preview(name, split, limit)

        return HuggingFacePreviewResponse(
            name=preview["name"],
            split=preview["split"],
            columns=preview["columns"],
            rows=preview["rows"],
            total_count=preview["total_count"]
        )
    except ValueError as e:
        if "not found" in str(e).lower():
            raise HTTPException(status_code=404, detail=str(e))
        elif "authentication" in str(e).lower() or "token" in str(e).lower():
            raise HTTPException(status_code=401, detail=str(e))
        else:
            raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to preview dataset: {str(e)}")


@router.post("/api/datasets/huggingface/import", response_model=DatasetResponse)
def import_huggingface_dataset(
    request: HuggingFaceImportRequest,
    db: Session = Depends(get_db)
):
    """Import a Hugging Face dataset into the system.

    Args:
        request: Import parameters including project_id, dataset_name, split, etc.

    Returns:
        Created Dataset object
    """
    _check_huggingface_enabled()

    from backend.dataset import HuggingFaceImporter

    # Validate project exists
    project = db.query(Project).filter(Project.id == request.project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    try:
        importer = HuggingFaceImporter(db)
        dataset = importer.import_dataset(
            project_id=request.project_id,
            dataset_name=request.dataset_name,
            split=request.split,
            display_name=request.display_name,
            row_limit=request.row_limit,
            columns=request.columns,
            add_row_id=request.add_row_id
        )

        row_count = importer.get_row_count(dataset.sqlite_table_name)

        return DatasetResponse(
            id=dataset.id,
            project_id=dataset.project_id,
            name=dataset.name,
            source_file_name=dataset.source_file_name,
            sqlite_table_name=dataset.sqlite_table_name,
            created_at=dataset.created_at,
            row_count=row_count
        )
    except ValueError as e:
        if "not found" in str(e).lower():
            raise HTTPException(status_code=404, detail=str(e))
        elif "authentication" in str(e).lower() or "token" in str(e).lower():
            raise HTTPException(status_code=401, detail=str(e))
        else:
            raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to import dataset: {str(e)}")


# ========== Excel Dataset Import Endpoints ==========

@router.post("/api/datasets/import", response_model=DatasetResponse)
async def import_dataset(
    project_id: int = Form(...),
    dataset_name: str = Form(...),
    range_name: str = Form("DSRange"),
    add_row_id: str = Form("false"),
    replace_dataset_id: Optional[int] = Form(None),
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Import dataset from Excel file.

    Specification: docs/req.txt section 4.6.2
    Phase 2

    Args:
        add_row_id: "true" to add RowID column as first column (starting from 1)
        replace_dataset_id: If provided, replace existing dataset (keep same ID)
    """
    # Validate file type
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(
            status_code=400,
            detail="Only Excel files (.xlsx, .xls) are supported"
        )

    # Parse boolean from form string
    add_row_id_bool = add_row_id.lower() in ("true", "1", "yes")

    # Save uploaded file to temporary location
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
        content = await file.read()
        tmp_file.write(content)
        tmp_file_path = tmp_file.name

    try:
        # Import dataset
        importer = DatasetImporter(db)
        dataset = importer.import_from_excel(
            project_id=project_id,
            file_path=tmp_file_path,
            dataset_name=dataset_name,
            range_name=range_name,
            add_row_id=add_row_id_bool,
            replace_dataset_id=replace_dataset_id
        )

        # Get row count
        row_count = importer._get_row_count(dataset.sqlite_table_name)

        return DatasetResponse(
            id=dataset.id,
            project_id=dataset.project_id,
            name=dataset.name,
            source_file_name=dataset.source_file_name,
            sqlite_table_name=dataset.sqlite_table_name,
            created_at=dataset.created_at,
            row_count=row_count
        )

    except HTTPException:
        raise  # Re-raise HTTPException as-is
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Import failed: {str(e)}")
    finally:
        # Clean up temporary file
        if os.path.exists(tmp_file_path):
            os.unlink(tmp_file_path)


@router.get("/api/datasets/{dataset_id}", response_model=DatasetResponse)
def get_dataset(dataset_id: int, db: Session = Depends(get_db)):
    """Get dataset details.

    Specification: docs/req.txt section 4.6
    Phase 2
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    importer = DatasetImporter(db)
    row_count = importer._get_row_count(dataset.sqlite_table_name)

    # Get all associated project IDs
    associations = db.query(ProjectDataset.project_id).filter(
        ProjectDataset.dataset_id == dataset.id
    ).all()
    associated_project_ids = [a[0] for a in associations]
    all_project_ids = list(set([dataset.project_id] + associated_project_ids))

    return DatasetResponse(
        id=dataset.id,
        project_id=dataset.project_id,
        name=dataset.name,
        source_file_name=dataset.source_file_name,
        sqlite_table_name=dataset.sqlite_table_name,
        created_at=dataset.created_at,
        row_count=row_count,
        project_ids=all_project_ids
    )


@router.get("/api/datasets/{dataset_id}/preview", response_model=DatasetPreviewResponse)
def preview_dataset(
    dataset_id: int,
    limit: int = 10,
    db: Session = Depends(get_db)
):
    """Get preview of dataset.

    Specification: docs/req.txt section 4.6 (implied)
    Phase 2
    """
    try:
        importer = DatasetImporter(db)
        preview = importer.get_dataset_preview(dataset_id, limit)

        return DatasetPreviewResponse(**preview)

    except HTTPException:
        raise  # Re-raise HTTPException as-is
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Preview failed: {str(e)}")


@router.get("/api/datasets/{dataset_id}/columns", response_model=List[str])
def get_dataset_columns(
    dataset_id: int,
    db: Session = Depends(get_db)
):
    """Get column names of a dataset.

    Args:
        dataset_id: Dataset ID

    Returns:
        List of column names
    """
    try:
        importer = DatasetImporter(db)
        # Get preview with limit=1 to get columns efficiently
        preview = importer.get_dataset_preview(dataset_id, limit=1)
        return preview.get("columns", [])

    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to get columns: {str(e)}")


class ColumnRenameRequest(BaseModel):
    """Request model for column rename."""
    column_mapping: Dict[str, str]


@router.put("/api/datasets/{dataset_id}/columns")
def rename_dataset_columns(
    dataset_id: int,
    request: ColumnRenameRequest,
    db: Session = Depends(get_db)
):
    """Rename columns in a dataset.

    Args:
        dataset_id: Dataset ID
        request: Column mapping {old_name: new_name, ...}

    Returns:
        Updated column list
    """
    try:
        importer = DatasetImporter(db)
        new_columns = importer.rename_columns(dataset_id, request.column_mapping)
        return {
            "success": True,
            "columns": new_columns
        }

    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to rename columns: {str(e)}")


class ColumnRestructureRequest(BaseModel):
    """Request model for column restructure (add/delete/reorder)."""
    columns: List[str]  # New ordered list of column names
    renames: Optional[Dict[str, str]] = None  # {old_name: new_name}


@router.put("/api/datasets/{dataset_id}/columns/restructure")
def restructure_dataset_columns(
    dataset_id: int,
    request: ColumnRestructureRequest,
    db: Session = Depends(get_db)
):
    """Restructure columns in a dataset: add, delete, reorder, rename.

    Args:
        dataset_id: Dataset ID
        request: Column restructure request
            - columns: Ordered list of final column names
            - renames: Optional mapping of old names to new names

    Returns:
        Updated column list
    """
    try:
        importer = DatasetImporter(db)
        new_columns = importer.restructure_columns(
            dataset_id,
            request.columns,
            request.renames
        )
        return {
            "success": True,
            "columns": new_columns
        }

    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to restructure columns: {str(e)}")


@router.get("/api/datasets/{dataset_id}/download")
def download_dataset(
    dataset_id: int,
    db: Session = Depends(get_db)
):
    """Download dataset as CSV file.

    Args:
        dataset_id: Dataset ID

    Returns:
        CSV file with UTF-8 BOM encoding
    """
    from fastapi.responses import StreamingResponse
    import io
    import csv

    # Get dataset
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    try:
        # Get all data from the dataset
        importer = DatasetImporter(db)
        preview = importer.get_dataset_preview(dataset_id, limit=0)  # 0 = no limit

        columns = preview.get("columns", [])
        rows = preview.get("rows", [])

        # Create CSV content with UTF-8 BOM
        output = io.StringIO()
        output.write('\ufeff')  # UTF-8 BOM for Excel compatibility

        writer = csv.writer(output)
        writer.writerow(columns)  # Header
        for row in rows:
            writer.writerow([row.get(col, '') for col in columns])

        csv_content = output.getvalue()
        output.close()

        # Return as streaming response
        # RFC 5987: Encode filename for non-ASCII characters
        from urllib.parse import quote
        ascii_name = ''.join(c if c.isascii() else '_' for c in dataset.name)
        utf8_name = quote(dataset.name)

        return StreamingResponse(
            iter([csv_content.encode('utf-8')]),
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f"attachment; filename=\"{ascii_name}.csv\"; filename*=UTF-8''{utf8_name}.csv"
            }
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Download failed: {str(e)}")


@router.delete("/api/datasets/{dataset_id}")
def delete_dataset(dataset_id: int, db: Session = Depends(get_db)):
    """Delete dataset.

    Specification: docs/req.txt section 4.6.1
    Phase 2
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    # Drop the SQLite table
    from sqlalchemy import text
    try:
        db.execute(text(f'DROP TABLE IF EXISTS "{dataset.sqlite_table_name}"'))
    except Exception:
        pass  # Continue even if table drop fails

    # Delete dataset record
    db.delete(dataset)
    db.commit()

    return {"success": True, "message": f"Dataset {dataset_id} deleted"}


# ========== Extended Import Endpoints ==========

class ImportFromJobRequest(BaseModel):
    """Request model for importing from job results."""
    job_id: int
    project_id: int
    dataset_name: Optional[str] = None
    target_dataset_id: Optional[int] = None
    add_row_id: bool = False


@router.post("/api/datasets/import/append", response_model=DatasetResponse)
async def append_excel_to_dataset(
    target_dataset_id: int = Form(...),
    range_name: str = Form("DSRange"),
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Append data from Excel file to existing dataset."""
    # Validate file type
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(
            status_code=400,
            detail="Only Excel files (.xlsx, .xls) are supported"
        )

    # Get target dataset
    dataset = db.query(Dataset).filter(Dataset.id == target_dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Target dataset not found")

    # Save uploaded file
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
        content = await file.read()
        tmp_file.write(content)
        tmp_file_path = tmp_file.name

    try:
        importer = DatasetImporter(db)

        # Extract data from Excel
        import openpyxl
        workbook = openpyxl.load_workbook(tmp_file_path, data_only=True)

        if range_name not in workbook.defined_names:
            raise HTTPException(status_code=400, detail=f"Named range '{range_name}' not found in workbook")

        data_rows = importer._extract_data_from_range(workbook, range_name)
        if not data_rows or len(data_rows) < 2:
            raise HTTPException(status_code=400, detail="No data found in named range")

        header = data_rows[0]
        data_rows = data_rows[1:]

        # Get existing columns
        pragma_sql = f'PRAGMA table_info("{dataset.sqlite_table_name}")'
        result = db.execute(text(pragma_sql))
        existing_columns = [row[1] for row in result if row[1] != "id"]

        # Sanitize and map columns
        columns = [importer._sanitize_column_name(col) for col in header]

        # Insert data
        for row in data_rows:
            row_data = row[:len(columns)]
            while len(row_data) < len(columns):
                row_data.append("")

            # Build insert for existing columns only
            insert_cols = []
            insert_vals = {}
            for i, col in enumerate(columns):
                if col in existing_columns:
                    insert_cols.append(col)
                    insert_vals[f"param{i}"] = row_data[i]

            if insert_cols:
                placeholders = ", ".join([f":param{i}" for i, col in enumerate(columns) if col in existing_columns])
                column_list = ", ".join([f'"{col}"' for col in insert_cols])
                insert_sql = f'INSERT INTO "{dataset.sqlite_table_name}" ({column_list}) VALUES ({placeholders})'
                db.execute(text(insert_sql), insert_vals)

        db.commit()

        row_count = importer._get_row_count(dataset.sqlite_table_name)

        return DatasetResponse(
            id=dataset.id,
            project_id=dataset.project_id,
            name=dataset.name,
            source_file_name=dataset.source_file_name,
            sqlite_table_name=dataset.sqlite_table_name,
            created_at=dataset.created_at,
            row_count=row_count
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Append failed: {str(e)}")
    finally:
        if os.path.exists(tmp_file_path):
            os.unlink(tmp_file_path)


@router.post("/api/datasets/import/csv", response_model=DatasetResponse)
async def import_csv_dataset(
    project_id: int = Form(...),
    encoding: str = Form("utf-8"),
    delimiter: str = Form(","),
    quotechar: str = Form('"'),
    has_header: str = Form("1"),
    dataset_name: Optional[str] = Form(None),
    target_dataset_id: Optional[int] = Form(None),
    add_row_id: str = Form("false"),
    replace_dataset_id: Optional[int] = Form(None),
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Import dataset from CSV file.

    Args:
        add_row_id: "true" to add RowID column as first column (starting from 1)
        replace_dataset_id: If provided, replace existing dataset (keep same ID)
    """
    # Parse boolean from form string
    add_row_id_bool = add_row_id.lower() in ("true", "1", "yes")

    # Validate inputs
    if not dataset_name and not target_dataset_id and not replace_dataset_id:
        raise HTTPException(status_code=400, detail="Either dataset_name, target_dataset_id, or replace_dataset_id is required")

    # Read file content
    content = await file.read()

    # Decode with specified encoding
    try:
        text_content = content.decode(encoding)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to decode file with encoding {encoding}: {str(e)}")

    # Parse CSV
    try:
        reader = csv.reader(
            io.StringIO(text_content),
            delimiter=delimiter,
            quotechar=quotechar if quotechar else '"'
        )
        rows = list(reader)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse CSV: {str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="CSV file is empty")

    # Handle header
    if has_header == "1":
        header = rows[0]
        data_rows = rows[1:]
    else:
        # Generate column names
        col_count = len(rows[0]) if rows else 0
        header = [f"col_{i+1}" for i in range(col_count)]
        data_rows = rows

    if not data_rows:
        raise HTTPException(status_code=400, detail="No data rows found")

    # Add RowID column if requested
    if add_row_id_bool:
        header = ["RowID"] + header
        for i, row in enumerate(data_rows, start=1):
            row.insert(0, str(i))

    importer = DatasetImporter(db)

    try:
        if replace_dataset_id:
            # Replace existing dataset
            dataset = importer.replace_dataset(
                replace_dataset_id,
                header,
                data_rows,
                file.filename or "import.csv"
            )

        elif target_dataset_id:
            # Append to existing dataset
            dataset = db.query(Dataset).filter(Dataset.id == target_dataset_id).first()
            if not dataset:
                raise HTTPException(status_code=404, detail="Target dataset not found")

            # Get existing columns
            pragma_sql = f'PRAGMA table_info("{dataset.sqlite_table_name}")'
            result = db.execute(text(pragma_sql))
            existing_columns = [row[1] for row in result if row[1] != "id"]

            # Sanitize columns
            columns = [importer._sanitize_column_name(col) for col in header]

            # Insert data
            for row in data_rows:
                row_data = row[:len(columns)]
                while len(row_data) < len(columns):
                    row_data.append("")

                insert_cols = []
                insert_vals = {}
                for i, col in enumerate(columns):
                    if col in existing_columns:
                        insert_cols.append(col)
                        insert_vals[f"param{i}"] = row_data[i]

                if insert_cols:
                    placeholders = ", ".join([f":param{i}" for i, col in enumerate(columns) if col in existing_columns])
                    column_list = ", ".join([f'"{col}"' for col in insert_cols])
                    insert_sql = f'INSERT INTO "{dataset.sqlite_table_name}" ({column_list}) VALUES ({placeholders})'
                    db.execute(text(insert_sql), insert_vals)

            db.commit()

        else:
            # Create new dataset
            timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S_%f")  # Add microseconds for uniqueness
            table_name = f"Dataset_PJ{project_id}_{timestamp}"

            dataset = Dataset(
                project_id=project_id,
                name=dataset_name,
                source_file_name=file.filename or "import.csv",
                sqlite_table_name=table_name
            )
            db.add(dataset)
            db.flush()

            # Create table and insert data
            importer._create_and_populate_table(table_name, header, data_rows)
            db.commit()
            db.refresh(dataset)

        row_count = importer._get_row_count(dataset.sqlite_table_name)

        return DatasetResponse(
            id=dataset.id,
            project_id=dataset.project_id,
            name=dataset.name,
            source_file_name=dataset.source_file_name,
            sqlite_table_name=dataset.sqlite_table_name,
            created_at=dataset.created_at,
            row_count=row_count
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Import failed: {str(e)}")


@router.post("/api/datasets/import/from-job", response_model=DatasetResponse)
async def import_from_job(
    request: ImportFromJobRequest,
    db: Session = Depends(get_db)
):
    """Import dataset from job execution results.

    IMPORTANT: Header and data order consistency
    - First try job.merged_csv_output (most reliable)
    - If csv_header exists: use csv_header + csv_output (same source, safe)
    - If csv_header missing: regenerate BOTH from fields (ignore csv_output)
    """
    # Get job
    job = db.query(Job).filter(Job.id == request.job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    all_rows = []
    header = None

    # First, try to use job.merged_csv_output if available (most reliable)
    if job.merged_csv_output:
        lines = [line.strip() for line in job.merged_csv_output.strip().split("\n") if line.strip()]
        if lines:
            # First line is header
            header = lines[0].split(",")
            # Remaining lines are data
            for line in lines[1:]:
                try:
                    reader = csv.reader(io.StringIO(line))
                    row = next(reader)
                    all_rows.append(row)
                except:
                    all_rows.append([line])
    else:
        # Fall back to building from job_items
        job_items = db.query(JobItem).filter(
            JobItem.job_id == request.job_id,
            JobItem.status == "done"
        ).all()

        if not job_items:
            raise HTTPException(status_code=400, detail="No completed job items found")

        for item in job_items:
            if not item.parsed_response:
                continue

            try:
                parsed = json.loads(item.parsed_response)
            except:
                continue

            csv_output = parsed.get("csv_output", "")
            csv_header = parsed.get("csv_header", "")
            fields = parsed.get("fields", {})

            # SAFETY: Ensure header and data come from the same source
            if csv_header:
                # csv_header exists: use csv_header + csv_output (same source, order matches)
                if header is None:
                    header = csv_header.split(",")
                if csv_output:
                    for line in csv_output.strip().split("\n"):
                        if line.strip():
                            try:
                                reader = csv.reader(io.StringIO(line))
                                row = next(reader)
                                all_rows.append(row)
                            except:
                                all_rows.append([line])
            elif fields:
                # csv_header missing: generate BOTH from fields (ignore csv_output)
                if header is None:
                    header = list(fields.keys())
                # Always use fields.values() in the order of header keys
                row = [str(fields.get(h, "")) for h in header]
                all_rows.append(row)

    if not all_rows:
        raise HTTPException(status_code=400, detail="No CSV data found in job results")

    # Generate header if not found (last resort)
    if header is None:
        col_count = len(all_rows[0]) if all_rows else 1
        header = [f"col_{i+1}" for i in range(col_count)]

    importer = DatasetImporter(db)

    try:
        if request.target_dataset_id:
            # Append to existing
            dataset = db.query(Dataset).filter(Dataset.id == request.target_dataset_id).first()
            if not dataset:
                raise HTTPException(status_code=404, detail="Target dataset not found")

            # Get existing columns
            pragma_sql = f'PRAGMA table_info("{dataset.sqlite_table_name}")'
            result = db.execute(text(pragma_sql))
            existing_columns = [row[1] for row in result if row[1] != "id"]

            # Add RowID if requested (append mode)
            if request.add_row_id:
                existing_count = importer._get_row_count(dataset.sqlite_table_name)
                # Add RowID column if it doesn't exist
                if "RowID" not in existing_columns:
                    alter_sql = f'ALTER TABLE "{dataset.sqlite_table_name}" ADD COLUMN "RowID" INTEGER'
                    db.execute(text(alter_sql))
                    existing_columns.append("RowID")
                # Add RowID to header and data
                header = ["RowID"] + header
                for i, row in enumerate(all_rows, start=existing_count + 1):
                    row.insert(0, i)

            columns = [importer._sanitize_column_name(col) for col in header]

            for row in all_rows:
                row_data = row[:len(columns)]
                while len(row_data) < len(columns):
                    row_data.append("")

                insert_cols = []
                insert_vals = {}
                for i, col in enumerate(columns):
                    if col in existing_columns:
                        insert_cols.append(col)
                        insert_vals[f"param{i}"] = row_data[i]

                if insert_cols:
                    placeholders = ", ".join([f":param{i}" for i, col in enumerate(columns) if col in existing_columns])
                    column_list = ", ".join([f'"{col}"' for col in insert_cols])
                    insert_sql = f'INSERT INTO "{dataset.sqlite_table_name}" ({column_list}) VALUES ({placeholders})'
                    db.execute(text(insert_sql), insert_vals)

            db.commit()

        else:
            # Create new
            if not request.dataset_name:
                raise HTTPException(status_code=400, detail="dataset_name is required for new dataset")

            timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
            table_name = f"Dataset_PJ{request.project_id}_{timestamp}"

            dataset = Dataset(
                project_id=request.project_id,
                name=request.dataset_name,
                source_file_name=f"job_{request.job_id}_results",
                sqlite_table_name=table_name
            )
            db.add(dataset)
            db.flush()

            # Add RowID if requested
            if request.add_row_id:
                header = ["RowID"] + header
                for i, row in enumerate(all_rows, start=1):
                    row.insert(0, i)

            importer._create_and_populate_table(table_name, header, all_rows)
            db.commit()
            db.refresh(dataset)

        row_count = importer._get_row_count(dataset.sqlite_table_name)

        return DatasetResponse(
            id=dataset.id,
            project_id=dataset.project_id,
            name=dataset.name,
            source_file_name=dataset.source_file_name,
            sqlite_table_name=dataset.sqlite_table_name,
            created_at=dataset.created_at,
            row_count=row_count
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Import failed: {str(e)}")


# ========== Dataset-Project Association Endpoints ==========

class DatasetProjectsRequest(BaseModel):
    """Request model for updating dataset project associations."""
    project_ids: List[int]


class ProjectDatasetInfo(BaseModel):
    """Project info for dataset association."""
    id: int
    name: str
    is_owner: bool = False


@router.get("/api/datasets/{dataset_id}/projects", response_model=List[ProjectDatasetInfo])
def get_dataset_projects(dataset_id: int, db: Session = Depends(get_db)):
    """Get list of projects associated with a dataset.

    Returns all projects that can access this dataset (owner + associated).
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    # Get owner project
    owner = db.query(Project).filter(Project.id == dataset.project_id).first()

    # Get associated projects
    associations = db.query(ProjectDataset).filter(
        ProjectDataset.dataset_id == dataset_id
    ).all()
    associated_project_ids = [a.project_id for a in associations]

    # Build result with all projects
    all_project_ids = list(set([dataset.project_id] + associated_project_ids))
    projects = db.query(Project).filter(Project.id.in_(all_project_ids)).all()

    result = []
    for project in projects:
        result.append(ProjectDatasetInfo(
            id=project.id,
            name=project.name,
            is_owner=(project.id == dataset.project_id)
        ))

    # Sort: owner first, then by name
    result.sort(key=lambda p: (not p.is_owner, p.name))
    return result


@router.put("/api/datasets/{dataset_id}/projects")
def update_dataset_projects(
    dataset_id: int,
    request: DatasetProjectsRequest,
    db: Session = Depends(get_db)
):
    """Update the list of projects associated with a dataset.

    This replaces all current associations. The owner project is always included.
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    # Validate all project IDs exist
    for project_id in request.project_ids:
        project = db.query(Project).filter(Project.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail=f"Project {project_id} not found")

    # Remove all existing associations (except owner relationship in datasets table)
    db.query(ProjectDataset).filter(ProjectDataset.dataset_id == dataset_id).delete()

    # Add new associations (excluding owner project as it's already in datasets.project_id)
    for project_id in request.project_ids:
        if project_id != dataset.project_id:  # Don't duplicate owner
            association = ProjectDataset(
                project_id=project_id,
                dataset_id=dataset_id
            )
            db.add(association)

    db.commit()

    # Return updated project list
    return get_dataset_projects(dataset_id, db)


@router.post("/api/datasets/{dataset_id}/projects/{project_id}")
def add_dataset_project(
    dataset_id: int,
    project_id: int,
    db: Session = Depends(get_db)
):
    """Add a single project association to a dataset."""
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Check if already associated (including owner)
    if project_id == dataset.project_id:
        return {"success": True, "message": "Project is already the owner"}

    existing = db.query(ProjectDataset).filter(
        ProjectDataset.dataset_id == dataset_id,
        ProjectDataset.project_id == project_id
    ).first()

    if existing:
        return {"success": True, "message": "Association already exists"}

    # Create new association
    association = ProjectDataset(
        project_id=project_id,
        dataset_id=dataset_id
    )
    db.add(association)
    db.commit()

    return {"success": True, "message": f"Project {project_id} added to dataset {dataset_id}"}


@router.delete("/api/datasets/{dataset_id}/projects/{project_id}")
def remove_dataset_project(
    dataset_id: int,
    project_id: int,
    db: Session = Depends(get_db)
):
    """Remove a single project association from a dataset.

    Cannot remove the owner project.
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    # Cannot remove owner
    if project_id == dataset.project_id:
        raise HTTPException(status_code=400, detail="Cannot remove owner project")

    # Find and delete association
    association = db.query(ProjectDataset).filter(
        ProjectDataset.dataset_id == dataset_id,
        ProjectDataset.project_id == project_id
    ).first()

    if not association:
        raise HTTPException(status_code=404, detail="Association not found")

    db.delete(association)
    db.commit()

    return {"success": True, "message": f"Project {project_id} removed from dataset {dataset_id}"}


# ========== Row CRUD Endpoints ==========

class RowData(BaseModel):
    """Request model for row data."""
    data: Dict[str, Any]


class RowPreviewResponse(BaseModel):
    """Dataset preview response with rowid."""
    dataset_id: int
    name: str
    columns: List[str]
    rows: List[Dict[str, Any]]  # Each row includes 'rowid' field
    total_count: int


@router.get("/api/datasets/{dataset_id}/rows", response_model=RowPreviewResponse)
def get_dataset_rows(
    dataset_id: int,
    limit: int = 10,
    offset: int = 0,
    db: Session = Depends(get_db)
):
    """Get dataset rows with rowid for editing.

    Returns rows with SQLite rowid for identification.
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    table_name = dataset.sqlite_table_name
    if not table_name:
        raise HTTPException(status_code=400, detail="Dataset has no table")

    try:
        # Get column names
        col_result = db.execute(text(f'PRAGMA table_info("{table_name}")'))
        columns = [row[1] for row in col_result]

        # Get total count
        count_result = db.execute(text(f'SELECT COUNT(*) FROM "{table_name}"'))
        total_count = count_result.scalar()

        # Get rows with rowid
        cols_sql = ', '.join([f'"{c}"' for c in columns])
        if limit > 0:
            sql = f'SELECT rowid, {cols_sql} FROM "{table_name}" LIMIT :limit OFFSET :offset'
            result = db.execute(text(sql), {"limit": limit, "offset": offset})
        else:
            sql = f'SELECT rowid, {cols_sql} FROM "{table_name}"'
            result = db.execute(text(sql))

        rows = []
        for row in result:
            row_dict = {"rowid": row[0]}  # First element is rowid
            for i, col in enumerate(columns):
                row_dict[col] = row[i + 1]  # Data starts at index 1
            rows.append(row_dict)

        return RowPreviewResponse(
            dataset_id=dataset_id,
            name=dataset.name,
            columns=columns,
            rows=rows,
            total_count=total_count
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to get rows: {str(e)}")


@router.post("/api/datasets/{dataset_id}/rows")
def add_dataset_row(
    dataset_id: int,
    request: RowData,
    db: Session = Depends(get_db)
):
    """Add a new row to the dataset.

    Args:
        dataset_id: Dataset ID
        request: Row data as {column: value, ...}

    Returns:
        Created row with rowid
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    table_name = dataset.sqlite_table_name
    if not table_name:
        raise HTTPException(status_code=400, detail="Dataset has no table")

    try:
        # Get valid column names
        col_result = db.execute(text(f'PRAGMA table_info("{table_name}")'))
        valid_columns = [row[1] for row in col_result]

        # Filter and prepare data
        insert_cols = []
        insert_vals = {}
        param_idx = 0
        for col, val in request.data.items():
            if col in valid_columns:
                insert_cols.append(col)
                insert_vals[f"p{param_idx}"] = val
                param_idx += 1

        if not insert_cols:
            raise HTTPException(status_code=400, detail="No valid columns provided")

        # Build and execute INSERT
        column_list = ', '.join([f'"{c}"' for c in insert_cols])
        placeholders = ', '.join([f":p{i}" for i in range(len(insert_cols))])
        insert_sql = f'INSERT INTO "{table_name}" ({column_list}) VALUES ({placeholders})'
        db.execute(text(insert_sql), insert_vals)
        db.commit()

        # Get the inserted rowid
        rowid_result = db.execute(text("SELECT last_insert_rowid()"))
        new_rowid = rowid_result.scalar()

        return {
            "success": True,
            "rowid": new_rowid,
            "message": f"Row added with rowid {new_rowid}"
        }

    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to add row: {str(e)}")


@router.put("/api/datasets/{dataset_id}/rows/{rowid}")
def update_dataset_row(
    dataset_id: int,
    rowid: int,
    request: RowData,
    db: Session = Depends(get_db)
):
    """Update an existing row in the dataset.

    Args:
        dataset_id: Dataset ID
        rowid: SQLite rowid of the row to update
        request: Updated row data as {column: value, ...}

    Returns:
        Success status
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    table_name = dataset.sqlite_table_name
    if not table_name:
        raise HTTPException(status_code=400, detail="Dataset has no table")

    try:
        # Verify row exists
        check_sql = f'SELECT rowid FROM "{table_name}" WHERE rowid = :rowid'
        result = db.execute(text(check_sql), {"rowid": rowid})
        if not result.fetchone():
            raise HTTPException(status_code=404, detail=f"Row with rowid {rowid} not found")

        # Get valid column names
        col_result = db.execute(text(f'PRAGMA table_info("{table_name}")'))
        valid_columns = [row[1] for row in col_result]

        # Build UPDATE statement
        set_clauses = []
        params = {"rowid": rowid}
        param_idx = 0
        for col, val in request.data.items():
            if col in valid_columns:
                set_clauses.append(f'"{col}" = :p{param_idx}')
                params[f"p{param_idx}"] = val
                param_idx += 1

        if not set_clauses:
            raise HTTPException(status_code=400, detail="No valid columns provided")

        update_sql = f'UPDATE "{table_name}" SET {", ".join(set_clauses)} WHERE rowid = :rowid'
        db.execute(text(update_sql), params)
        db.commit()

        return {
            "success": True,
            "rowid": rowid,
            "message": f"Row {rowid} updated"
        }

    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to update row: {str(e)}")


@router.delete("/api/datasets/{dataset_id}/rows/{rowid}")
def delete_dataset_row(
    dataset_id: int,
    rowid: int,
    db: Session = Depends(get_db)
):
    """Delete a row from the dataset.

    Args:
        dataset_id: Dataset ID
        rowid: SQLite rowid of the row to delete

    Returns:
        Success status
    """
    dataset = db.query(Dataset).filter(Dataset.id == dataset_id).first()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    table_name = dataset.sqlite_table_name
    if not table_name:
        raise HTTPException(status_code=400, detail="Dataset has no table")

    try:
        # Verify row exists
        check_sql = f'SELECT rowid FROM "{table_name}" WHERE rowid = :rowid'
        result = db.execute(text(check_sql), {"rowid": rowid})
        if not result.fetchone():
            raise HTTPException(status_code=404, detail=f"Row with rowid {rowid} not found")

        # Delete row
        delete_sql = f'DELETE FROM "{table_name}" WHERE rowid = :rowid'
        db.execute(text(delete_sql), {"rowid": rowid})
        db.commit()

        return {
            "success": True,
            "rowid": rowid,
            "message": f"Row {rowid} deleted"
        }

    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Failed to delete row: {str(e)}")
